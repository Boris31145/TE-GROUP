"""Create editable legal drafts from the same clauses used by the local site."""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend.legal_content import CONTRACT, SPECIFICATION, POLICY, POLICY_REGISTER, ACCOUNT_REQUEST, CONTRACT_REQUEST

OUT=ROOT.parent/'output/legal'
OUT.mkdir(parents=True,exist_ok=True)

def base(title):
    d=Document();s=d.sections[0]
    # The runtime template includes a title border; legal documents use whitespace.
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.85)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(11.5 if name=='Normal' else 19 if name=='Title' else 13)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=1.10
    d.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    d.styles['Heading 1'].paragraph_format.keep_with_next=True
    d.styles['Normal'].paragraph_format.widow_control=True
    d.add_heading(title,0)
    f=s.footer.paragraphs[0];f.alignment=2
    r=f.add_run();r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    d.core_properties.author='TE GROUP';d.core_properties.title=title
    d.core_properties.subject='Автомобильное направление'
    return d

def blocks(d,sections):
    for title,items in sections:
        d.add_heading(title,1)
        for p in items:d.add_paragraph(p)

def fields(d,rows):
    t=d.add_table(rows=0,cols=2);t.autofit=False
    widths=[Inches(1.8),Inches(5.0)]
    for column,width in zip(t.columns,widths):column.width=width
    for label,value in rows:
        cells=t.add_row().cells
        for c,w in zip(cells,widths):c.width=w
        cells[0].text=label;cells[1].text=value
        for r in cells[0].paragraphs[0].runs:r.bold=True
        for c in cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
                for r in p.runs:r.font.size=Pt(11)
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D5D5D5');b.append(e)
            pr.append(b)
        trpr=cells[0]._tc.getparent().get_or_add_trPr();keep=OxmlElement('w:cantSplit');trpr.append(keep)
    return t

def signatures(d):
    d.add_paragraph('Продавец: ____________________ / [ФИО, должность]')
    d.add_paragraph('Покупатель: __________________ / [ФИО]')
    d.add_paragraph('Дата подписания: [дата]')

d=base('Проект договора продажи автомобиля с доставкой')
d.add_paragraph('Договор № [номер]     [место подписания]     [дата]')
d.add_paragraph('ОсОО «ТЕ ГРУПП», ИНН 00506202410612, адрес: Кыргызстан, Бишкек, ул. Киевская, 77, офис 603, в лице [ФИО и должность], действующего на основании [устав / доверенность с реквизитами], далее Продавец, и [ФИО покупателя], адрес [адрес], далее Покупатель, заключили настоящий договор.')
blocks(d,CONTRACT)
d.add_heading('11 Реквизиты и подписи',1)
fields(d,[('Продавец','ОсОО «ТЕ ГРУПП»; ИНН 00506202410612; Кыргызстан, Бишкек, ул. Киевская, 77, офис 603. Регистрационные сведения и полномочия: [заполнить].'),('Банк и счет','[Получатель, банк, страна, счет, банковский код, необходимые корреспондентские реквизиты, валюта].'),('Контакты Продавца','Email для деловой связи: oootegroup@gmail.com. Проверенный канал значимых сообщений и приема претензий: [заполнить].'),('Покупатель','[ФИО, адрес, телефон, email; идентификационные сведения только в необходимом объеме].')])
signatures(d)
d.add_page_break();d.add_heading('Приложение 1 Спецификация автомобиля и расчет',1)
d.add_paragraph('К договору № [номер] от [дата]. Составляется и подписывается до оплаты.')
fields(d,SPECIFICATION[:6])
d.add_page_break();d.add_heading('Приложение 1 Продолжение спецификации',1)
fields(d,SPECIFICATION[6:]);signatures(d)
d.add_page_break();d.add_heading('Приложение 2 Акт передачи автомобиля',1)
d.add_paragraph('К договору № [номер] от [дата]. Место и дата фактической передачи: [значение].')
fields(d,[('Автомобиль','Марка / модель [значение]; VIN [значение]; пробег на момент передачи [значение].'),('Получатель','Покупатель / представитель [ФИО]; основание полномочий [реквизиты].'),('Комплектность','Ключи, документы и принадлежности: [перечень, количество, оригинал / копия].'),('Состояние','Видимые повреждения и расхождения со Спецификацией: [перечень либо прямо «не обнаружены при осмотре»]. Фото: [приложения].'),('Таможенный статус','[Статус на момент передачи, сведения о дальнейших действиях и ответственных лицах].'),('Замечания и действия','[Требования покупателя, согласованные действия, сроки].'),('Порядок возврата','Письменная памятка выдана: [дата, версия, способ]. Адрес приема и доступные способы возврата: [значение].')])
d.add_paragraph('Акт фиксирует фактическую передачу, комплектность и результат доступного осмотра. Он не является отказом от требований о скрытых недостатках, недостоверных сведениях или иных законных прав. При выявленном несоответствии факт и согласованные действия указываются выше.')
signatures(d);d.save(OUT/'TE_GROUP_Car_Sale_Contract.docx')

d=base('Единая политика конфиденциальности и обработки персональных данных')
d.add_paragraph('Автомобильное направление TE GROUP')
blocks(d,POLICY)
d.add_page_break();d.add_heading('Приложение 1 Реестр инфраструктуры и сроков',1)
fields(d,POLICY_REGISTER)
d.add_paragraph('Реестр содержит конкретные сведения до начала соответствующей обработки. При отсутствии определенного основания или необходимой процедуры соответствующая обработка не запускается. Политика и реестр утверждаются уполномоченным лицом Оператора.')
d.add_page_break();d.add_heading('Приложение 2 Заявления пользователя',1)
d.add_heading('Создание личного кабинета',2)
for p in ACCOUNT_REQUEST:d.add_paragraph(p)
d.add_paragraph('Подтверждение: [непредустановленный флажок / иной доказуемый способ]. Время, версия и идентификатор пользователя: [запись системы].')
d.add_heading('Подготовка договора по выбранной брони',2)
for p in CONTRACT_REQUEST:d.add_paragraph(p)
d.add_paragraph('Бронь: [номер]. Подтверждение: [непредустановленный флажок / иной доказуемый способ]. Время, версия текста и редакция сведений: [запись системы].')
d.add_heading('Дополнительные цели',2)
d.add_paragraph('Настоящие заявления подтверждают инициативу пользователя и не являются согласиями на дополнительные цели. Для рекламы, публикации или иной цели, требующей согласия, предоставляется отдельный конкретный текст и отдельное необязательное подтверждение. Такие цели в автомобильном сервисе данной редакции не предусмотрены.')
d.add_heading('Согласие на поручение обработки при необходимости',2)
d.add_paragraph('Я предоставляю ОсОО «ТЕ ГРУПП», ИНН 00506202410612, адрес: Кыргызстан, Бишкек, ул. Киевская, 77, офис 603, отдельное согласие на поручение обработки [точный перечень данных] лицу [полное наименование обработчика, адрес, страна] исключительно для [конкретная цель]. Допускаются [необходимые действия]. Способ обработки: [автоматизированный / смешанный].')
d.add_paragraph('Согласие действует [определенный срок либо событие прекращения]. Отзыв направляется через кабинет либо по утвержденному адресу для запросов Оператора. Оператор прекращает обработку на основании этого согласия в предусмотренном законом порядке; самостоятельные законные основания оцениваются отдельно. Согласие не разрешает рекламу, публикацию данных, произвольных новых получателей и не отменяет требований локализации и трансграничной передачи.')
d.add_paragraph('Этот текст предоставляется пользователю отдельно от политики и заявления о создании кабинета только после заполнения всех индивидуальных полей и проверки необходимости согласия. Подтверждение самостоятельное, непредустановленное и фиксируется с сохранением текста, версии, времени и идентификатора пользователя. Наличие образца в приложении не означает получения согласия.')
d.save(OUT/'TE_GROUP_Privacy_and_Data_Policy.docx')
print('\n'.join(str(p) for p in OUT.glob('*.docx')))
